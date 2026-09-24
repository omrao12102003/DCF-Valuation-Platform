from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def _write_dataframe(
    ws,
    df: pd.DataFrame,
    start_row: int = 1,
    start_col: int = 1,
) -> int:
    if df.empty:
        ws.cell(start_row, start_col, "No data")
        return start_row + 1

    for col_index, column in enumerate(df.columns, start_col):
        cell = ws.cell(start_row, col_index, str(column))
        cell.font = Font(bold=True)

    for row_offset, (_, row) in enumerate(df.iterrows(), 1):
        for col_offset, value in enumerate(row, start_col):
            if pd.isna(value):
                value = None
            ws.cell(start_row + row_offset, col_offset, value)

    return start_row + len(df) + 2


def _format_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))

        ws.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12),
            32,
        )


def _add_title(ws, title: str) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="left")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)


def export_valuation_workbook(
    output_path: str | Path,
    *,
    ticker: str,
    historical: pd.DataFrame,
    forecast: pd.DataFrame | None = None,
    ufcf: pd.DataFrame | None = None,
    dcf: Any | None = None,
    sensitivity: pd.DataFrame | None = None,
    reverse_dcf: pd.DataFrame | None = None,
    comparables: pd.DataFrame | None = None,
    roic: pd.DataFrame | None = None,
    valuation_summary: pd.DataFrame | None = None,
) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    sheets: list[tuple[str, pd.DataFrame | None]] = [
        ("Historical", historical),
        ("Forecast", forecast),
        ("UFCF", ufcf),
        ("Sensitivity", sensitivity),
        ("Reverse DCF", reverse_dcf),
        ("Comparables", comparables),
        ("ROIC", roic),
        ("Valuation", valuation_summary),
    ]

    for name, dataframe in sheets:
        if dataframe is None:
            continue

        ws = workbook.create_sheet(name)
        _add_title(ws, f"{ticker} - {name}")
        ws["A3"] = "Development model output"
        ws["A3"].font = Font(italic=True)

        if isinstance(dataframe, pd.DataFrame):
            _write_dataframe(ws, dataframe, start_row=5)

        _format_sheet(ws)

    if dcf is not None:
        ws = workbook.create_sheet("DCF")
        _add_title(ws, f"{ticker} - DCF Valuation")

        values: list[tuple[str, Any]] = []

        for field in (
            "enterprise_value",
            "equity_value",
            "value_per_share",
            "terminal_value",
            "terminal_value_percentage",
            "net_debt",
            "shares_outstanding",
            "wacc",
            "terminal_growth",
        ):
            if hasattr(dcf, field):
                values.append((field, getattr(dcf, field)))

        ws["A3"] = "Metric"
        ws["B3"] = "Value"

        for row_number, (metric, value) in enumerate(values, 4):
            ws.cell(row_number, 1, metric)
            ws.cell(row_number, 2, value)

        _format_sheet(ws)

    workbook.properties.title = f"{ticker} DCF Valuation Platform"
    workbook.properties.subject = "Integrated financial valuation model"
    workbook.properties.creator = "DCF Valuation Platform"

    workbook.save(output)

    return output


def validate_workbook(
    path: str | Path,
    *,
    required_sheets: list[str],
) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=False)

    missing = [
        sheet
        for sheet in required_sheets
        if sheet not in workbook.sheetnames
    ]

    result = {
        "valid": not missing,
        "sheets": workbook.sheetnames,
        "missing_sheets": missing,
    }

    workbook.close()
    return result
