from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.finance.excel_export import (
    export_valuation_workbook,
    validate_workbook,
)


def test_export_creates_workbook(tmp_path: Path):
    historical = pd.DataFrame(
        {
            "fiscal_year": [2024, 2025],
            "revenue": [100.0, 110.0],
            "ebit": [20.0, 23.0],
        }
    )

    forecast = pd.DataFrame(
        {
            "fiscal_year": [2026, 2027],
            "revenue": [120.0, 132.0],
            "ebit": [25.0, 28.0],
        }
    )

    output = tmp_path / "valuation.xlsx"

    result = export_valuation_workbook(
        output,
        ticker="MSFT",
        historical=historical,
        forecast=forecast,
        valuation_summary=pd.DataFrame(
            {
                "Method": ["DCF", "Trading Comparables"],
                "Value Per Share": [250.0, 240.0],
            }
        ),
    )

    assert result.exists()

    workbook = load_workbook(result, read_only=True)
    assert "Historical" in workbook.sheetnames
    assert "Forecast" in workbook.sheetnames
    assert "Valuation" in workbook.sheetnames
    workbook.close()


def test_workbook_validation(tmp_path: Path):
    output = tmp_path / "valuation.xlsx"

    export_valuation_workbook(
        output,
        ticker="AAPL",
        historical=pd.DataFrame(
            {
                "fiscal_year": [2025],
                "revenue": [100.0],
            }
        ),
    )

    result = validate_workbook(
        output,
        required_sheets=["Historical"],
    )

    assert result["valid"] is True
    assert result["missing_sheets"] == []


def test_workbook_validation_detects_missing_sheet(tmp_path: Path):
    output = tmp_path / "valuation.xlsx"

    export_valuation_workbook(
        output,
        ticker="AAPL",
        historical=pd.DataFrame(
            {
                "fiscal_year": [2025],
                "revenue": [100.0],
            }
        ),
    )

    result = validate_workbook(
        output,
        required_sheets=["Historical", "DCF"],
    )

    assert result["valid"] is False
    assert "DCF" in result["missing_sheets"]


def test_export_preserves_numeric_values(tmp_path: Path):
    output = tmp_path / "valuation.xlsx"

    export_valuation_workbook(
        output,
        ticker="NVDA",
        historical=pd.DataFrame(
            {
                "fiscal_year": [2025],
                "revenue": [202841.0],
                "ebitda": [94755.0],
            }
        ),
    )

    workbook = load_workbook(output, data_only=False)
    ws = workbook["Historical"]

    values = [
        cell.value
        for row in ws.iter_rows()
        for cell in row
    ]

    assert 202841.0 in values
    assert 94755.0 in values

    workbook.close()
